# -*- coding: utf-8 -*-
"""
生成汇报 PPT（16:9）。
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image
import os

BASE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(BASE, 'figures')

# ---------- 配色 ----------
NAVY   = RGBColor(0x1F, 0x38, 0x64)
BLUE   = RGBColor(0x2E, 0x75, 0xB6)
RED    = RGBColor(0xD6, 0x45, 0x45)
GREEN  = RGBColor(0x3E, 0x9B, 0x6C)
ORANGE = RGBColor(0xE8, 0xA3, 0x3D)
GRAY   = RGBColor(0x59, 0x59, 0x59)
LIGHT  = RGBColor(0xEE, 0xF3, 0xF8)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x33, 0x33, 0x33)
MIDGRAY= RGBColor(0x80, 0x80, 0x80)

FONT = 'Microsoft YaHei'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

SW, SH = prs.slide_width, prs.slide_height


def set_run_font(run, size=18, bold=False, color=DARK, italic=False, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    # 设置东亚字体
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', font)


def add_text(slide, x, y, w, h, lines, size=18, bold=False, color=DARK,
             align=PP_ALIGN.LEFT, font=FONT, line_spacing=1.12, anchor=MSO_ANCHOR.TOP):
    """lines: 字符串或 [(text, size, bold, color), ...] 列表"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    if isinstance(lines, str):
        lines = [(lines, size, bold, color)]
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if isinstance(item, str):
            item = (item, size, bold, color)
        text, sz, bd, cl = item
        run = p.add_run()
        run.text = text
        set_run_font(run, sz, bd, cl, font=font)
    return tb


def add_rect(slide, x, y, w, h, fill=LIGHT, line=NAVY, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
             text=None, text_color=NAVY, text_size=14, bold=True, radius=0.06):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    sp.line.color.rgb = line
    sp.line.width = Pt(line_w)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    sp.shadow.inherit = False
    if text:
        tf = sp.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.03)
        tf.margin_right = Inches(0.03)
        tf.margin_top = Inches(0.02)
        tf.margin_bottom = Inches(0.02)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        lines = text.split('\n')
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = ln
            set_run_font(r, text_size, bold, text_color)
    return sp


def add_arrow(slide, x1, y1, x2, y2, color=BLUE, width=2.0):
    conn = slide.shapes.add_connector(1, x1, y1, x2, y2)  # 1 = straight
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    # 箭头
    line_elem = conn.line._get_or_add_ln()
    tail = line_elem.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    line_elem.append(tail)
    conn.shadow.inherit = False
    return conn


def add_image_fit(slide, path, x, y, max_w, max_h, align='center'):
    """按比例缩放图片，放入 (x,y,max_w,max_h) 框内，水平对齐。"""
    im = Image.open(path)
    iw, ih = im.size
    ar = iw / ih
    box_ar = max_w / max_h
    if ar >= box_ar:
        w = max_w
        h = int(max_w / ar)
    else:
        h = max_h
        w = int(max_h * ar)
    if align == 'center':
        x = x + (max_w - w) // 2
    y = y + (max_h - h) // 2
    slide.shapes.add_picture(path, x, y, w, h)
    return (w, h)


def add_title_bar(slide, title, subtitle=None, num=None):
    """顶部标题栏。"""
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(0.95))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    bar.shadow.inherit = False
    # 左侧竖条
    acc = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), Inches(0.95))
    acc.fill.solid(); acc.fill.fore_color.rgb = ORANGE
    acc.line.fill.background(); acc.shadow.inherit = False
    prefix = f'{num}  ' if num else ''
    add_text(slide, Inches(0.45), Inches(0.14), Inches(12.5), Inches(0.7),
             f'{prefix}{title}', size=26, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_text(slide, Inches(0.5), Inches(1.0), Inches(12.5), Inches(0.4),
                 subtitle, size=13, bold=False, color=MIDGRAY)


def add_footer(slide, idx, total):
    add_text(slide, Inches(12.2), Inches(7.05), Inches(1.0), Inches(0.4),
             f'{idx} / {total}', size=11, color=MIDGRAY, align=PP_ALIGN.RIGHT)


TOTAL = 10


# ======================================================================
# 1. 封面
# ======================================================================
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background(); bg.shadow.inherit = False
# 装饰横条
for i, (y, c) in enumerate([(1.0, ORANGE), (1.15, BLUE), (1.30, GREEN)]):
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.4+i*0.18), Inches(2.4), Inches(0.06))
    b.fill.solid(); b.fill.fore_color.rgb = c; b.line.fill.background(); b.shadow.inherit = False
add_text(s, Inches(1.0), Inches(2.9), Inches(11.3), Inches(1.2),
         '基于轻量迭代优化的自主泊车轨迹规划', size=40, bold=True, color=WHITE)
add_text(s, Inches(1.0), Inches(4.05), Inches(11.3), Inches(0.6),
         'Optimization-Based Autonomous Parking Trajectory Planning (LIOM)', size=20, bold=False, color=RGBColor(0xBF, 0xD3, 0xE8))
add_text(s, Inches(1.0), Inches(5.0), Inches(11.3), Inches(0.5),
         '研究进展汇报', size=22, bold=False, color=RGBColor(0xDD, 0xE7, 0xF2))
add_text(s, Inches(1.0), Inches(6.4), Inches(11.3), Inches(0.6),
         [('汇报内容：项目五个工作模块的完成情况与效果展示', 14, False, RGBColor(0xA8, 0xC0, 0xDA))])


# ======================================================================
# 2. 项目概述
# ======================================================================
s = prs.slides.add_slide(BLANK)
add_title_bar(s, '项目概述与研究内容', num='01')
add_text(s, Inches(0.6), Inches(1.12), Inches(12.1), Inches(1.25),
         [('项目目标：', 17, True, NAVY),
          ('实现并工程化论文《Optimization-Based Trajectory Planning for Autonomous Parking With '
           'Irregularly Placed Obstacles》(Bai Li et al., IEEE TITS 2022) 中的 LIOM 算法，'
           '在 ROS2 / Nav2 平台上完成不规则障碍环境下的自主泊车轨迹规划。', 17, False, DARK)],
         line_spacing=1.18)
# 技术路线
add_text(s, Inches(0.6), Inches(2.45), Inches(6.0), Inches(0.4),
         '技术路线（三段式）', size=16, bold=True, color=BLUE)
steps = ['Hybrid A* 粗路径', '安全走廊生成 (SFC)', 'IPOPT 轨迹优化']
x = Inches(0.7)
for i, st in enumerate(steps):
    add_rect(s, x, Inches(2.88), Inches(3.4), Inches(0.8), fill=NAVY, line=NAVY,
             text=st, text_color=WHITE, text_size=16)
    if i < 2:
        add_arrow(s, x + Inches(3.4), Inches(3.28), x + Inches(4.1), Inches(3.28), color=ORANGE, width=2.5)
    x += Inches(3.7)
add_text(s, Inches(0.7), Inches(3.8), Inches(11.0), Inches(0.4),
         '粗路径提供运动学可行初值 → 迭代扩张碰撞安全走廊 → 软约束 OCP 求解出平滑轨迹', size=13, color=MIDGRAY)
# 五个工作点
add_text(s, Inches(0.6), Inches(4.32), Inches(6.0), Inches(0.4),
         '本次汇报的五个工作模块', size=16, bold=True, color=BLUE)
mods = [
    ('common_math', '2D 几何数学库（Apollo 移植）', BLUE),
    ('car_description', '车辆 URDF 模型 + 圆盘碰撞表示', GREEN),
    ('liom_local_planner', '核心：Hybrid A* + 走廊 + IPOPT 优化', RED),
    ('planner', 'Nav2 全局规划器插件封装', ORANGE),
    ('simulator', '轻量仿真环境 + 人机交互', RGBColor(0x8A, 0x5A, 0xB5)),
]
y = Inches(4.76)
for i, (name, desc, c) in enumerate(mods):
    add_rect(s, Inches(0.7), y + Inches(0.44*i), Inches(3.1), Inches(0.4),
             fill=LIGHT, line=c, text=name, text_color=NAVY, text_size=13, radius=0.5)
    add_text(s, Inches(4.0), y + Inches(0.44*i), Inches(8.6), Inches(0.4),
             desc, size=14, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
add_footer(s, 2, TOTAL)


# ======================================================================
# 3. 总体架构
# ======================================================================
s = prs.slides.add_slide(BLANK)
add_title_bar(s, '系统总体架构', num='02')
# 第一层：仿真与交互
add_text(s, Inches(0.5), Inches(1.1), Inches(3.0), Inches(0.35), '仿真与交互层', size=14, bold=True, color=GRAY)
add_rect(s, Inches(0.7), Inches(1.5), Inches(3.4), Inches(0.9), fill=LIGHT, line=GREEN,
         text='simulator\n时钟 / 里程计 / 遥操作 / 目标发布', text_size=13)
add_rect(s, Inches(4.6), Inches(1.5), Inches(3.2), Inches(0.9), fill=LIGHT, line=GREEN,
         text='RViz2 交互界面\n点选目标 · 轨迹显示', text_size=13)
add_arrow(s, Inches(4.1), Inches(1.95), Inches(4.6), Inches(1.95), color=GREEN, width=2.0)
# 第二层：Nav2 框架
add_text(s, Inches(0.5), Inches(2.65), Inches(3.0), Inches(0.35), 'Nav2 规划框架', size=14, bold=True, color=GRAY)
add_rect(s, Inches(0.7), Inches(3.0), Inches(3.4), Inches(0.85), fill=NAVY, line=NAVY,
         text='planner_server\nComputePathToPose', text_color=WHITE, text_size=13)
add_rect(s, Inches(4.6), Inches(3.0), Inches(4.2), Inches(0.85), fill=NAVY, line=NAVY,
         text='CustomPlanner 插件 (planner)\nnav2_core::GlobalPlanner', text_color=WHITE, text_size=13)
add_arrow(s, Inches(4.1), Inches(3.42), Inches(4.6), Inches(3.42), color=BLUE, width=2.0)
add_arrow(s, Inches(2.4), Inches(3.85), Inches(2.4), Inches(4.35), color=ORANGE, width=2.5)
# 第三层：核心算法
add_rect(s, Inches(0.7), Inches(4.35), Inches(8.1), Inches(1.6), fill=RGBColor(0xFD, 0xEF, 0xEF),
         line=RED, line_w=1.8,
         text='liom_local_planner — 核心算法\nHybrid A* 粗路径  →  走廊生成 (SFC)  →  IPOPT 轨迹优化 (ADOL-C)',
         text_color=RED, text_size=15)
# 第四层：基础层
add_text(s, Inches(9.2), Inches(1.1), Inches(3.6), Inches(0.35), '基础支撑层', size=14, bold=True, color=GRAY)
add_rect(s, Inches(9.2), Inches(1.5), Inches(3.6), Inches(1.3), fill=LIGHT, line=BLUE,
         text='common_math\n2D 几何库\nVec2d/Pose/Box/Polygon/样条', text_size=13)
add_rect(s, Inches(9.2), Inches(3.0), Inches(3.6), Inches(1.3), fill=LIGHT, line=GREEN,
         text='car_description\n车辆 URDF 模型\n+ 圆盘碰撞表示', text_size=13)
add_rect(s, Inches(9.2), Inches(4.55), Inches(3.6), Inches(1.4), fill=LIGHT, line=RGBColor(0x8A, 0x5A, 0xB5),
         text='costmap / OMPL / IPOPT\n外部依赖\n(地图 · 解析曲线 · 求解器)', text_size=13)
add_footer(s, 3, TOTAL)


# ======================================================================
# 4. 工作点1 common_math
# ======================================================================
s = prs.slides.add_slide(BLANK)
add_title_bar(s, '工作点 1 — common_math 几何数学库', subtitle='基础层：为上层规划提供 2D 几何与数值计算支撑')
items = [
    ('几何类型', 'Vec2d / Pose / AABox2d / Box2d / Polygon2d / LineSegment2d'),
    ('碰撞与距离', '点/线段/盒/多边形之间的重叠判断、距离计算、凸包与 IoU'),
    ('数值工具', '角度归一化与连续化、Clamp、线性插值 lerp / Interpolate1d、LinSpaced'),
    ('样条插值', '三次 B 样条控制点求解、求值与求导（供 planner 做轨迹稠密化）'),
    ('来源与定位', '由 Apollo 开源几何库移植，作为项目底层依赖被各模块复用'),
]
y = Inches(1.35)
for i, (h, d) in enumerate(items):
    add_rect(s, Inches(0.7), y + Inches(0.88*i), Inches(2.3), Inches(0.72), fill=NAVY, line=NAVY,
             text=h, text_color=WHITE, text_size=14)
    add_text(s, Inches(3.3), y + Inches(0.88*i), Inches(9.3), Inches(0.72),
             d, size=15, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(0.7), Inches(6.0), Inches(12.0), Inches(0.7),
         [('说明：', 15, True, ORANGE),
          ('该模块为纯基础设施，不直接产生可视化效果；其几何原语与样条工具是后续“轨迹优化 + 插值输出”的底层支撑。', 15, False, DARK)])
add_footer(s, 4, TOTAL)


# ======================================================================
# 5. 工作点2 car_description
# ======================================================================
s = prs.slides.add_slide(BLANK)
add_title_bar(s, '工作点 2 — car_description 车辆建模', subtitle='后轴中心坐标系 + 圆盘碰撞表示，尺寸与运动学参数严格一致')
add_text(s, Inches(0.6), Inches(1.2), Inches(4.6), Inches(5.4),
         [('建模内容', 17, True, NAVY),
          ('• 后轮轴中心为基准的车辆 URDF（车身 + 四轮 + 标记）', 15, False, DARK),
          ('• 车长 4.689 m / 车宽 1.942 m / 轴距 2.80 m', 15, False, DARK),
          ('• 前轮可转向关节 (±0.61 rad)，含传动配置', 15, False, DARK),
          ('', 8, False, DARK),
          ('圆盘碰撞模型', 17, True, NAVY),
          ('• 用 2 个圆盘覆盖车辆矩形轮廓', 15, False, DARK),
          ('• 圆盘半径 = 0.5·√((L/n)² + W²) ≈ 1.52 m', 15, False, DARK),
          ('• 把碰撞检测简化为“圆盘中心是否在障碍物内”', 15, False, DARK),
          ('', 8, False, DARK),
          ('作用', 17, True, NAVY),
          ('• RViz 中显示车辆本体', 15, False, DARK),
          ('• 参数与 PlannerConfig 运动学模型一一对应', 15, False, DARK)],
         line_spacing=1.12)
add_image_fit(s, os.path.join(FIG, 'vehicle.png'), Inches(5.4), Inches(1.25), Inches(7.3), Inches(5.6))
add_footer(s, 5, TOTAL)


# ======================================================================
# 6. 工作点3 核心算法 + 泊车效果
# ======================================================================
s = prs.slides.add_slide(BLANK)
add_title_bar(s, '工作点 3 — liom_local_planner 核心算法', subtitle='项目核心：Hybrid A* 粗路径 + 安全走廊 + IPOPT 轨迹优化')
# 左列：6 个子模块
add_text(s, Inches(0.5), Inches(1.15), Inches(3.6), Inches(0.35), '核心子模块', size=15, bold=True, color=BLUE)
subs = [
    ('VehicleModel', '车辆运动学 + 圆盘碰撞'),
    ('Environment', 'R-tree 碰撞检测 + 走廊生成'),
    ('PathPlanner', 'Hybrid A* 粗路径(Reeds-Shepp)'),
    ('LightweightProblem', 'IPOPT OCP(软约束 + ADOL-C)'),
    ('LiomLocalPlanner', '迭代走廊优化主循环(Alg.3)'),
    ('visualization', 'RViz 走廊/轨迹可视化'),
]
for i, (h, d) in enumerate(subs):
    add_rect(s, Inches(0.5), Inches(1.55) + Inches(0.86*i), Inches(2.15), Inches(0.72),
             fill=LIGHT, line=NAVY, text=h, text_color=NAVY, text_size=12)
    add_text(s, Inches(2.8), Inches(1.55) + Inches(0.86*i), Inches(3.1), Inches(0.72),
             d, size=11.5, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
# 右侧：效果图
add_image_fit(s, os.path.join(FIG, 'parking.png'), Inches(5.9), Inches(1.15), Inches(7.1), Inches(5.7))
add_footer(s, 6, TOTAL)


# ======================================================================
# 7. 工作点3 迭代走廊优化框架
# ======================================================================
s = prs.slides.add_slide(BLANK)
add_title_bar(s, '迭代走廊优化框架（论文 Alg. 3）', subtitle='核心贡献：坏初值下逐步重建走廊并热启动求解，直至收敛')
add_text(s, Inches(0.5), Inches(1.15), Inches(12.3), Inches(0.75),
         [('思路：', 15, True, NAVY),
          ('建走廊 → 解盒约束 OCP → 计算不可行度 ψ → 用当前解重建走廊并热启动再解，循环至 ψ < ε。'
           '相比“只建一次走廊解一次”的 STC 方法，能稳定从坏初值恢复。', 15, False, DARK)], line_spacing=1.2)
add_image_fit(s, os.path.join(FIG, 'convergence.png'), Inches(0.5), Inches(2.0), Inches(12.4), Inches(3.3))
add_image_fit(s, os.path.join(FIG, 'corridor.png'), Inches(3.2), Inches(5.35), Inches(6.9), Inches(1.9))
add_footer(s, 7, TOTAL)


# ======================================================================
# 8. 工作点4 planner
# ======================================================================
s = prs.slides.add_slide(BLANK)
add_title_bar(s, '工作点 4 — planner Nav2 插件封装', subtitle='将 LIOM 算法接入 Nav2 全局规划框架')
add_text(s, Inches(0.6), Inches(1.3), Inches(12.1), Inches(4.6),
         [('插件实现', 17, True, NAVY),
          ('• CustomPlanner 继承 nav2_core::GlobalPlanner，通过 pluginlib 注册', 15, False, DARK),
          ('• createPlan() 调用 LiomLocalPlanner::Plan() 完成全局规划', 15, False, DARK),
          ('• 以 planner/CustomPlanner 插件形式在 planner.yaml 中启用', 15, False, DARK),
          ('', 8, False, DARK),
          ('轨迹稠密化', 17, True, NAVY),
          ('• 优化结果为稀疏状态点，用三次 B 样条按时间插值', 15, False, DARK),
          ('• 解缠航向、由几何反解 v/φ、差分求 a/ω，输出 nav_msgs::Path', 15, False, DARK),
          ('', 8, False, DARK),
          ('工程整理', 17, True, NAVY),
          ('• 旧版 Hybrid A*(SmacPlannerHybrid 风格)代码已注释并标注“已弃用”', 15, False, DARK),
          ('• 统一改用 liom_local_planner，接口清晰、可维护', 15, False, DARK)],
         line_spacing=1.15)
# 右侧流水线
add_rect(s, Inches(7.2), Inches(1.6), Inches(5.4), Inches(0.8), fill=NAVY, line=NAVY,
         text='createPlan(start, goal)', text_color=WHITE, text_size=14)
add_rect(s, Inches(7.2), Inches(2.7), Inches(5.4), Inches(0.8), fill=RGBColor(0xFD, 0xEF, 0xEF), line=RED,
         text='LiomLocalPlanner::Plan()\n(粗路径 → 走廊 → IPOPT)', text_color=RED, text_size=13)
add_rect(s, Inches(7.2), Inches(3.8), Inches(5.4), Inches(0.8), fill=LIGHT, line=BLUE,
         text='三次 B 样条按时间插值\n稠密化轨迹', text_color=NAVY, text_size=13)
add_rect(s, Inches(7.2), Inches(4.9), Inches(5.4), Inches(0.8), fill=LIGHT, line=GREEN,
         text='nav_msgs::Path\n返回 Nav2 / RViz', text_color=GREEN, text_size=13)
for yy in (2.4, 3.5, 4.6):
    add_arrow(s, Inches(9.9), Inches(yy), Inches(9.9), Inches(yy+0.3), color=ORANGE, width=2.0)
add_footer(s, 8, TOTAL)


# ======================================================================
# 9. 工作点5 simulator
# ======================================================================
s = prs.slides.add_slide(BLANK)
add_title_bar(s, '工作点 5 — simulator 仿真环境', subtitle='轻量仿真闭环：时钟 / 里程计 / 遥操作 / 目标发布')
add_text(s, Inches(0.6), Inches(1.3), Inches(12.1), Inches(5.2),
         [('仿真节点', 17, True, NAVY),
          ('• clock_node — 发布 /clock 仿真时钟（10 ms 步进）', 15, False, DARK),
          ('• transform — 运动学机器人模型：订阅 /cmd_vel，积分里程计并广播 TF', 15, False, DARK),
          ('• telecontrol — 键盘遥操作（方向键发 /cmd_vel）', 15, False, DARK),
          ('• planner_bridge — RViz 点选目标，转成 Nav2 ComputePathToPose 请求', 15, False, DARK),
          ('', 8, False, DARK),
          ('与 Nav2 / RViz 集成', 17, True, NAVY),
          ('• 加载 car_description 车辆模型 + map_server 静态地图', 15, False, DARK),
          ('• 闭环：点选目标 → 规划 → 路径回显，用于验证规划效果', 15, False, DARK),
          ('• 支持 use_sim_time，与仿真时钟同步', 15, False, DARK)])
# 右侧闭环示意
cx = Inches(9.2)
add_rect(s, cx, Inches(1.5), Inches(3.4), Inches(0.7), fill=NAVY, line=NAVY,
         text='RViz 点选目标', text_color=WHITE, text_size=13)
add_rect(s, cx, Inches(2.5), Inches(3.4), Inches(0.7), fill=LIGHT, line=GREEN,
         text='planner_bridge\nComputePathToPose', text_color=NAVY, text_size=12)
add_rect(s, cx, Inches(3.5), Inches(3.4), Inches(0.7), fill=RGBColor(0xFD, 0xEF, 0xEF), line=RED,
         text='Nav2 规划器 (custom)', text_color=RED, text_size=13)
add_rect(s, cx, Inches(4.5), Inches(3.4), Inches(0.7), fill=LIGHT, line=BLUE,
         text='路径回显 + 车辆模型', text_color=NAVY, text_size=13)
add_rect(s, cx, Inches(5.5), Inches(3.4), Inches(0.7), fill=LIGHT, line=ORANGE,
         text='transform 里程计 / TF', text_color=NAVY, text_size=13)
for yy in (2.2, 3.2, 4.2, 5.2):
    add_arrow(s, cx + Inches(1.7), Inches(yy), cx + Inches(1.7), Inches(yy+0.3), color=ORANGE, width=2.0)
add_arrow(s, cx + Inches(3.4), Inches(1.85), cx + Inches(4.0), Inches(1.85), color=GREEN, width=2.0)
add_arrow(s, cx + Inches(4.0), Inches(1.85), cx + Inches(4.0), Inches(5.85), color=MIDGRAY, width=1.5)
add_arrow(s, cx + Inches(4.0), Inches(5.85), cx + Inches(3.4), Inches(5.85), color=ORANGE, width=2.0)
add_footer(s, 9, TOTAL)


# ======================================================================
# 10. 总结与下一步
# ======================================================================
s = prs.slides.add_slide(BLANK)
add_title_bar(s, '总结与下一步计划', num='03')
add_text(s, Inches(0.6), Inches(1.3), Inches(6.0), Inches(0.4), '已完成工作', size=18, bold=True, color=GREEN)
dones = [
    '完成 LIOM 算法完整实现（粗路径 + 走廊 + IPOPT 优化）',
    '补齐论文核心贡献 —— 迭代走廊优化框架（Alg. 3）',
    '修复时间步长 off-by-one、ADOL-C tape 瘦身等正确性/性能问题',
    '封装为 Nav2 全局规划器插件，接入 ROS2 平台',
    '搭建轻量仿真环境，支持 RViz 可视化验证',
]
y = Inches(1.75)
for i, d in enumerate(dones):
    add_text(s, Inches(0.6), y + Inches(0.52*i), Inches(6.3), Inches(0.5),
             f'✓  {d}', size=14, color=DARK)
add_text(s, Inches(7.2), Inches(1.3), Inches(5.5), Inches(0.4), '下一步计划', size=18, bold=True, color=ORANGE)
nexts = [
    '在更多复杂/不规则障碍场景下开展系统测试',
    '与基准方法（纯 Hybrid A* 等）做定量对比实验',
    '性能 profile 与进一步工程优化',
    '撰写论文正文与实验结果分析',
]
y = Inches(1.75)
for i, d in enumerate(nexts):
    add_text(s, Inches(7.2), y + Inches(0.52*i), Inches(5.6), Inches(0.5),
             f'▸  {d}', size=14, color=DARK)
add_footer(s, 10, TOTAL)


out = os.path.join(BASE, '汇报PPT.pptx')
prs.save(out)
print('saved:', out)
